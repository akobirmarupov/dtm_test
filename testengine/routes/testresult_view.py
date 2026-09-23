from __future__ import annotations

from io import BytesIO

from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView
from django.http import Http404, HttpResponse
from django.utils import timezone

from drf_spectacular.types import OpenApiTypes
from drf_spectacular.utils import OpenApiParameter, extend_schema

import logging

from testengine.filters import TestResultFilter
from testengine.models import TestResult
from testengine.routes.serializers import TestResultSerializer
from billing.entitlements import entitlements_for_request, history_cutoff
from common.i18n import resolve_language
from common.pagination import StandardResultsPagination
from common.permissions import HasFeature
from common.throttles import BurstUserRateThrottle


result_logger = logging.getLogger('testengine.result')

# Bitta faylga ko'chiriladigan maksimal qator. Cheksiz eksport bitta so'rovda
# butun bazani o'qib, xotirani yeb qo'yishi mumkin.
EXPORT_ROW_LIMIT = 5000

CanExportResults = HasFeature.named(
    'export_results',
    message="Natijalarni yuklab olish sizning tarifingizda mavjud emas.",
)


def serializer_context(request):
    return {'request': request, 'language': resolve_language(request)}


RESULT_SELECT_RELATED = ('session', 'session__subject', 'session__user')


def visible_results(request):
    """Foydalanuvchi ko'ra oladigan natijalar.

    Tarifda «Natijalar tarixi (kun)» qo'yilgan bo'lsa, undan eskisi
    ro'yxatga tushmaydi.
    """
    queryset = TestResult.objects.filter(
        session__user=request.user
    ).select_related(*RESULT_SELECT_RELATED).order_by('-created_at')

    cutoff = history_cutoff(entitlements_for_request(request))
    if cutoff is not None:
        queryset = queryset.filter(created_at__gte=cutoff)
    return queryset


class TestResultListAPIView(APIView):
    throttle_classes = [BurstUserRateThrottle]
    permission_classes = [IsAuthenticated]

    @extend_schema(responses=TestResultSerializer(many=True), tags=['TestResult'])
    def get(self, request):
        queryset = TestResultFilter(
            request.query_params, queryset=visible_results(request)
        ).qs

        paginator = StandardResultsPagination()
        page = paginator.paginate_queryset(queryset, request, view=self)
        serializer = TestResultSerializer(
            page, many=True, context=serializer_context(request)
        )

        return paginator.get_paginated_response(serializer.data)


class TestResultDetailAPIView(APIView):
    throttle_classes = [BurstUserRateThrottle]
    permission_classes = [IsAuthenticated]

    def get_object(self, pk, request):
        try:
            return visible_results(request).get(id=pk)
        except TestResult.DoesNotExist:
            raise Http404("Natija topilmadi.")

    @extend_schema(responses=TestResultSerializer, tags=['TestResult'])
    def get(self, request, pk):
        """Bitta test natijasi tafsiloti."""
        result = self.get_object(pk, request)
        return Response(
            TestResultSerializer(result, context=serializer_context(request)).data
        )


class MyTestResultsAPIView(APIView):
    throttle_classes = [BurstUserRateThrottle]
    permission_classes = [IsAuthenticated]

    @extend_schema(responses=TestResultSerializer(many=True), tags=['TestResult'])
    def get(self, request):
        queryset = TestResultFilter(
            request.query_params, queryset=visible_results(request)
        ).qs

        paginator = StandardResultsPagination()
        page = paginator.paginate_queryset(queryset, request, view=self)
        serializer = TestResultSerializer(
            page, many=True, context=serializer_context(request)
        )

        result_logger.info(
            "Natijalar ro'yxati ko'rildi: user_id=%s", request.user.id
        )

        return paginator.get_paginated_response(serializer.data)


class TestResultExportAPIView(APIView):
    """`GET /testengine/results/export/` — natijalarni fayl qilib beradi.

    `?type=xlsx` (standart) yoki `?type=pdf`. Fayl serverda saqlanmaydi,
    javobning o'zida qaytadi. Tarifdagi «Natijalar tarixi» chegarasi bu yerda
    ham amal qiladi: ko'rinmaydigan natija eksportga ham tushmaydi.
    """

    throttle_classes = [BurstUserRateThrottle]
    permission_classes = [IsAuthenticated, CanExportResults]

    COLUMNS = [
        ('Sana', 22),
        ('Fan', 22),
        ('Rejim', 14),
        ('Savollar', 10),
        ("To'g'ri", 10),
        ('Xato', 10),
        ('Javobsiz', 10),
        ('Ball', 10),
        ('Davomiyligi (daq.)', 18),
    ]

    @extend_schema(
        parameters=[
            OpenApiParameter(
                'type', OpenApiTypes.STR,
                description="`xlsx` (standart) yoki `pdf`. Nomi ataylab `format` "
                            "emas: `format` ni DRF o'zi kontent turi uchun ishlatadi.",
                enum=['xlsx', 'pdf'],
            ),
        ],
        responses={
            (200, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'):
                OpenApiTypes.BINARY,
        },
        tags=['TestResult'],
        description="Natijalar tarixini `.xlsx` yoki `.pdf` fayl qilib yuklab olish.",
    )
    def get(self, request):
        rows = self.build_rows(request)
        file_format = (request.query_params.get('type') or 'xlsx').lower()

        if file_format == 'pdf':
            content, content_type, extension = self.render_pdf(rows), 'application/pdf', 'pdf'
        elif file_format in ('xlsx', 'excel', ''):
            content = self.render_xlsx(rows)
            content_type = (
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            extension = 'xlsx'
        else:
            return Response(
                {
                    'detail': "Faqat `xlsx` yoki `pdf` formatida yuklab olish mumkin.",
                    'code': 'unsupported_format',
                },
                status=400,
            )

        filename = f"natijalar-{timezone.localdate():%Y-%m-%d}.{extension}"
        response = HttpResponse(content, content_type=content_type)
        response['Content-Disposition'] = f'attachment; filename="{filename}"'

        result_logger.info(
            'Natijalar eksport qilindi: user_id=%s format=%s qatorlar=%s',
            request.user.id, extension, len(rows),
        )
        return response

    def build_rows(self, request) -> list[list]:
        """Ikkala format ham shu qatorlardan chiziladi."""
        queryset = TestResultFilter(
            request.query_params, queryset=visible_results(request)
        ).qs[:EXPORT_ROW_LIMIT]

        rows = []
        for result in queryset:
            session = result.session
            rows.append([
                timezone.localtime(result.created_at).strftime('%Y-%m-%d %H:%M'),
                session.subject.name if session.subject_id else '—',
                session.get_mode_display(),
                session.question_count,
                result.correct_count,
                result.incorrect_count,
                result.unanswered_count,
                result.total_score,
                round(result.duration_seconds / 60, 1),
            ])
        return rows

    def render_xlsx(self, rows) -> bytes:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Natijalar'

        for index, (title, width) in enumerate(self.COLUMNS, start=1):
            cell = sheet.cell(row=1, column=index, value=title)
            cell.font = Font(bold=True)
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = 'A2'

        for row_number, values in enumerate(rows, start=2):
            for column, value in enumerate(values, start=1):
                sheet.cell(row=row_number, column=column, value=value)

        stream = BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    def render_pdf(self, rows) -> bytes:
        """Chop etishga tayyor jadval.

        Uzun ro'yxat sahifalarga o'zi bo'linadi, sarlavha har sahifada
        takrorlanadi (`repeatRows`).
        """
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        stream = BytesIO()
        document = SimpleDocTemplate(
            stream, pagesize=landscape(A4),
            leftMargin=12 * mm, rightMargin=12 * mm,
            topMargin=12 * mm, bottomMargin=12 * mm,
            title='Natijalar',
        )
        styles = getSampleStyleSheet()

        table_data = [[title for title, _ in self.COLUMNS]]
        table_data += [[str(value) for value in row] for row in rows]

        table = Table(table_data, repeatRows=1)
        table.setStyle(TableStyle([
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#EEF2F7')),
            ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#C9D2DD')),
            ('ALIGN', (3, 1), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))

        document.build([
            Paragraph('Test natijalari', styles['Title']),
            Paragraph(
                f"Yuklangan sana: {timezone.localtime():%Y-%m-%d %H:%M} · "
                f"Jami: {len(rows)} ta natija",
                styles['Normal'],
            ),
            Spacer(1, 8 * mm),
            table,
        ])
        return stream.getvalue()
